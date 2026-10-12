from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from pathlib import Path
from sqlalchemy.orm import Session
from sqlalchemy import func
import json
from collections import Counter
from openpyxl import Workbook

from app.database.connection import get_db
from app.models.user import User
from app.models.role import Role
from app.routes.auth import get_current_user
from app.models.tarot_reading import TarotReading
from app.models.palm_analysis import PalmAnalysis
from app.services.reports.report_generator import generate_palm_report


router = APIRouter(
    prefix="/admin",
    tags=["Administrator"]
)


def check_admin(current_user):
    if not current_user.role or current_user.role.role_name != "Administrator":
        raise HTTPException(
            status_code=403,
            detail="Administrator access required"
        )


# =========================================================
# USER MANAGEMENT
# =========================================================

@router.get("/users")
def get_all_users(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    check_admin(current_user)

    users = (
        db.query(User)
        .order_by(User.created_at.desc())
        .all()
    )

    return {
        "success": True,
        "total_users": len(users),
        "users": [
            {
                "user_id": user.user_id,
                "username": user.username,
                "email": user.email,
                "gender": user.gender,
                "role": (
                    user.role.role_name
                    if user.role
                    else "Unknown"
                ),
                "created_at": (
                    user.created_at.isoformat()
                    if user.created_at
                    else None
                )
            }
            for user in users
        ]
    }


@router.get("/users-by-role")
def get_users_by_role(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    check_admin(current_user)

    roles = db.query(Role).all()

    result = []

    for role in roles:
        count = (
            db.query(User)
            .filter(User.role_id == role.role_id)
            .count()
        )

        result.append({
            "role": role.role_name,
            "count": count
        })

    return {
        "success": True,
        "users_by_role": result
    }


@router.get("/recent-users")
def get_recent_users(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    check_admin(current_user)

    users = (
        db.query(User)
        .filter(User.created_at.isnot(None))
        .order_by(User.created_at.desc())
        .limit(5)
        .all()
    )

    return {
        "success": True,
        "recent_users": [
            {
                "user_id": user.user_id,
                "username": user.username,
                "email": user.email,
                "role": (
                    user.role.role_name
                    if user.role
                    else "Unknown"
                ),
                "created_at": (
                    user.created_at.isoformat()
                    if user.created_at
                    else None
                )
            }
            for user in users
        ]
    }

@router.get("/platform-analytics")
def get_platform_analytics(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    check_admin(current_user)

    # Total users
    total_users = db.query(User).count()

    # Total Palm analyses
    total_palm_analyses = db.query(PalmAnalysis).count()

    # Total Tarot readings
    total_tarot_readings = db.query(TarotReading).count()

    # Total combined activity
    total_analyses = (
        total_palm_analyses +
        total_tarot_readings
    )

    # Users by role
    roles = db.query(Role).all()

    users_by_role = []

    for role in roles:
        count = (
            db.query(User)
            .filter(User.role_id == role.role_id)
            .count()
        )

        users_by_role.append({
            "role": role.role_name,
            "count": count
        })

    # Platform activity trend
    palm_activity = (
        db.query(
            func.date(PalmAnalysis.created_at).label("date"),
            func.count(PalmAnalysis.analysis_id).label("count")
        )
        .group_by(func.date(PalmAnalysis.created_at))
        .all()
    )

    tarot_activity = (
        db.query(
            func.date(TarotReading.created_at).label("date"),
            func.count(TarotReading.reading_id).label("count")
        )
        .group_by(func.date(TarotReading.created_at))
        .all()
    )

    # Combine Palm and Tarot activity by date
    activity_by_date = {}

    for date, count in palm_activity:
        date_str = str(date)

        if date_str not in activity_by_date:
            activity_by_date[date_str] = {
                "date": date_str,
                "palm": 0,
                "tarot": 0,
                "total": 0
            }

        activity_by_date[date_str]["palm"] = count
        activity_by_date[date_str]["total"] += count

    for date, count in tarot_activity:
        date_str = str(date)

        if date_str not in activity_by_date:
            activity_by_date[date_str] = {
                "date": date_str,
                "palm": 0,
                "tarot": 0,
                "total": 0
            }

        activity_by_date[date_str]["tarot"] = count
        activity_by_date[date_str]["total"] += count

    activity_trend = sorted(
        activity_by_date.values(),
        key=lambda x: x["date"]
    )

    return {
        "success": True,

        "overview": {
            "total_users": total_users,
            "total_palm_analyses": total_palm_analyses,
            "total_tarot_readings": total_tarot_readings,
            "total_analyses": total_analyses
        },

        "analysis_distribution": {
            "palm_analysis": total_palm_analyses,
            "tarot_readings": total_tarot_readings
        },

        "users_by_role": users_by_role,

        "activity_trend": activity_trend
    }

# =========================================================
# READING STATISTICS
# =========================================================

@router.get("/reading-statistics")
def get_reading_statistics(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    check_admin(current_user)

    # =====================================================
    # TAROT SPREAD USAGE
    # =====================================================

    tarot_readings = (
        db.query(TarotReading)
        .order_by(TarotReading.created_at.desc())
        .all()
    )

    spread_counter = Counter()

    for reading in tarot_readings:
        if reading.spread_name:
            spread_counter[reading.spread_name] += 1

    tarot_spread_usage = [
        {
            "spread_name": spread_name,
            "count": count
        }
        for spread_name, count
        in spread_counter.most_common()
    ]

    # =====================================================
    # TAROT READING PATTERNS
    # =====================================================

    card_counts = [
        reading.card_count
        for reading in tarot_readings
        if reading.card_count is not None
    ]

    average_cards_per_reading = (
        round(sum(card_counts) / len(card_counts), 2)
        if card_counts
        else 0
    )

    card_count_distribution = Counter(card_counts)

    tarot_card_count_distribution = [
        {
            "card_count": card_count,
            "reading_count": count
        }
        for card_count, count
        in sorted(card_count_distribution.items())
    ]

    # =====================================================
    # PALM PATTERN STATISTICS
    # =====================================================

    palm_analyses = (
        db.query(PalmAnalysis)
        .order_by(PalmAnalysis.created_at.desc())
        .all()
    )

    palm_shape_counter = Counter()
    finger_structure_counter = Counter()
    palm_line_counter = Counter()
    palm_line_category_counter = Counter()

    for analysis in palm_analyses:

        if not analysis.analysis_result:
            continue

        try:
            result = json.loads(
                analysis.analysis_result
            )
        except (json.JSONDecodeError, TypeError):
            continue

        # -----------------------------------------------
        # Palm Shape
        # -----------------------------------------------

        palm_shape = result.get(
            "palm_shape", {}
        )

        if isinstance(palm_shape, dict):

            shape_name = palm_shape.get(
                "palm_shape"
            )

            if shape_name:
                palm_shape_counter[shape_name] += 1

        # -----------------------------------------------
        # Finger Structure
        # -----------------------------------------------

        finger_analysis = result.get(
            "finger_analysis", {}
        )

        if isinstance(finger_analysis, dict):

            finger_structure = finger_analysis.get(
                "finger_structure"
            )

            if finger_structure:
                finger_structure_counter[
                    finger_structure
                ] += 1

        # -----------------------------------------------
        # Palm Lines
        # -----------------------------------------------

        palm_lines = result.get(
            "palm_lines", {}
        )

        if isinstance(palm_lines, dict):

            lines = palm_lines.get(
                "lines", []
            )

            if isinstance(lines, list):

                for line in lines:

                    if not isinstance(line, dict):
                        continue

                    line_name = line.get("line")

                    if line_name:
                        palm_line_counter[
                            line_name
                        ] += 1

                    length_category = line.get(
                        "length_category"
                    )

                    if length_category:
                        palm_line_category_counter[
                            length_category
                        ] += 1

    # =====================================================
    # FORMAT PALM STATISTICS
    # =====================================================

    palm_shape_statistics = [
        {
            "palm_shape": shape,
            "count": count
        }
        for shape, count
        in palm_shape_counter.most_common()
    ]

    finger_structure_statistics = [
        {
            "finger_structure": structure,
            "count": count
        }
        for structure, count
        in finger_structure_counter.most_common()
    ]

    palm_line_statistics = [
        {
            "line": line,
            "count": count
        }
        for line, count
        in palm_line_counter.most_common()
    ]

    palm_line_category_statistics = [
        {
            "category": category,
            "count": count
        }
        for category, count
        in palm_line_category_counter.most_common()
    ]

    # =====================================================
    # READING ACTIVITY
    # =====================================================

    activity_counter = Counter()

    for reading in tarot_readings:

        if reading.created_at:
            date_key = reading.created_at.date().isoformat()
            activity_counter[date_key] += 1

    for analysis in palm_analyses:

        if analysis.created_at:
            date_key = analysis.created_at.date().isoformat()
            activity_counter[date_key] += 1

    reading_activity = [
        {
            "date": date,
            "count": count
        }
        for date, count
        in sorted(activity_counter.items())
    ]

    # =====================================================
    # FINAL RESPONSE
    # =====================================================

    return {
        "success": True,

        "tarot_statistics": {
            "spread_usage": tarot_spread_usage,
            "average_cards_per_reading":
                average_cards_per_reading,
            "card_count_distribution":
                tarot_card_count_distribution
        },

        "palm_statistics": {
            "palm_shape_usage":
                palm_shape_statistics,
            "finger_structure_usage":
                finger_structure_statistics,
            "palm_line_usage":
                palm_line_statistics,
            "palm_line_category_usage":
                palm_line_category_statistics
        },

        "reading_activity": reading_activity
    }

@router.get("/system-reports")
def get_system_reports(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    check_admin(current_user)

    try:
        # -----------------------------------------
        # PALMISTRY REPORTS
        # -----------------------------------------

        palm_analyses = (
            db.query(PalmAnalysis)
            .order_by(PalmAnalysis.created_at.desc())
            .all()
        )

        palm_reports = []

        for analysis in palm_analyses:

            result = {}

            if analysis.analysis_result:
                try:
                    result = json.loads(
                        analysis.analysis_result
                    )
                except (json.JSONDecodeError, TypeError):
                    result = {}

            palm_shape = ""
            finger_structure = ""

            palm_shape_data = result.get(
                "palm_shape", {}
            )

            if isinstance(palm_shape_data, dict):
                palm_shape = palm_shape_data.get(
                    "palm_shape",
                    ""
                )

            finger_data = result.get(
                "finger_analysis", {}
            )

            if isinstance(finger_data, dict):
                finger_structure = finger_data.get(
                    "finger_structure",
                    ""
                )

            palm_reports.append({
                "analysis_id": analysis.analysis_id,
                "user_id": analysis.user_id,
                "palm_shape": palm_shape,
                "finger_structure": finger_structure,
                "created_at": (
                    analysis.created_at.isoformat()
                    if analysis.created_at
                    else None
                )
            })

        # -----------------------------------------
        # TAROT REPORTS
        # -----------------------------------------

        tarot_readings = (
            db.query(TarotReading)
            .order_by(TarotReading.created_at.desc())
            .all()
        )

        tarot_reports = []

        for reading in tarot_readings:

            tarot_reports.append({
                "reading_id": reading.reading_id,
                "user_id": reading.user_id,
                "spread_name": reading.spread_name,
                "spread_type": reading.spread_type,
                "card_count": reading.card_count,
                "created_at": (
                    reading.created_at.isoformat()
                    if reading.created_at
                    else None
                )
            })

        # -----------------------------------------
        # PLATFORM ACTIVITY
        # -----------------------------------------

        activity_counter = Counter()

        for analysis in palm_analyses:

            if analysis.created_at:
                date = (
                    analysis.created_at
                    .date()
                    .isoformat()
                )

                activity_counter[date] += 1

        for reading in tarot_readings:

            if reading.created_at:
                date = (
                    reading.created_at
                    .date()
                    .isoformat()
                )

                activity_counter[date] += 1

        platform_activity = [
            {
                "date": date,
                "count": count
            }
            for date, count
            in sorted(activity_counter.items())
        ]

        # -----------------------------------------
        # SUMMARY
        # -----------------------------------------

        total_users = db.query(User).count()

        return {
            "success": True,

            "summary": {
                "total_users": total_users,
                "total_palm_reports": len(
                    palm_reports
                ),
                "total_tarot_reports": len(
                    tarot_reports
                ),
                "total_reports": (
                    len(palm_reports)
                    + len(tarot_reports)
                )
            },

            "palm_reports": palm_reports,

            "tarot_reports": tarot_reports,

            "platform_activity": platform_activity
        }

    except Exception as e:

        raise HTTPException(
            status_code=500,
            detail=(
                "Unable to generate system reports: "
                f"{str(e)}"
            )
        )

@router.get("/palm-report/{analysis_id}")
def get_admin_palm_report(
    analysis_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    check_admin(current_user)

    # Find the saved palm analysis
    analysis = (
        db.query(PalmAnalysis)
        .filter(PalmAnalysis.analysis_id == analysis_id)
        .first()
    )

    if not analysis:
        raise HTTPException(
            status_code=404,
            detail="Palm analysis not found."
        )

    # Convert saved JSON back into a Python dictionary
    try:
        palm_data = json.loads(
            analysis.analysis_result or "{}"
        )
    except (json.JSONDecodeError, TypeError):
        raise HTTPException(
            status_code=500,
            detail="Saved palm analysis data is invalid."
        )

    # Create reports folder
    reports_dir = Path("reports")
    reports_dir.mkdir(
        parents=True,
        exist_ok=True
    )

    # PDF file path
    output_path = (
        reports_dir /
        f"palm_analysis_{analysis_id}.pdf"
    )

    try:
        # Generate PDF using your existing generator
        generate_palm_report(
            palm_data,
            str(output_path)
        )

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Unable to generate Palm report: {str(e)}"
        )

    return FileResponse(
        path=str(output_path),
        media_type="application/pdf",
        filename=f"palm_analysis_{analysis_id}.pdf"
    )

@router.get("/system-reports/export-excel")
def export_system_reports_excel(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    check_admin(current_user)

    try:
        # -----------------------------------------
        # FETCH DATA
        # -----------------------------------------

        palm_analyses = (
            db.query(PalmAnalysis)
            .order_by(PalmAnalysis.created_at.desc())
            .all()
        )

        tarot_readings = (
            db.query(TarotReading)
            .order_by(TarotReading.created_at.desc())
            .all()
        )

        total_users = db.query(User).count()

        # -----------------------------------------
        # CREATE EXCEL WORKBOOK
        # -----------------------------------------

        workbook = Workbook()

        # -----------------------------------------
        # SUMMARY SHEET
        # -----------------------------------------

        summary_sheet = workbook.active
        summary_sheet.title = "Summary"

        summary_sheet.append(["System Report Summary"])
        summary_sheet.append([])

        summary_sheet.append([
            "Total Users",
            total_users
        ])

        summary_sheet.append([
            "Palm Reports",
            len(palm_analyses)
        ])

        summary_sheet.append([
            "Tarot Reports",
            len(tarot_readings)
        ])

        summary_sheet.append([
            "Total Reports",
            len(palm_analyses) + len(tarot_readings)
        ])

        # -----------------------------------------
        # PALMISTRY REPORTS SHEET
        # -----------------------------------------

        palm_sheet = workbook.create_sheet(
            "Palmistry Reports"
        )

        palm_sheet.append([
            "Analysis ID",
            "User ID",
            "Palm Shape",
            "Finger Structure",
            "Date"
        ])

        for analysis in palm_analyses:

            result = {}

            if analysis.analysis_result:
                try:
                    result = json.loads(
                        analysis.analysis_result
                    )
                except (
                    json.JSONDecodeError,
                    TypeError
                ):
                    result = {}

            palm_shape = ""
            finger_structure = ""

            palm_shape_data = result.get(
                "palm_shape",
                {}
            )

            if isinstance(palm_shape_data, dict):
                palm_shape = palm_shape_data.get(
                    "palm_shape",
                    ""
                )

            finger_data = result.get(
                "finger_analysis",
                {}
            )

            if isinstance(finger_data, dict):
                finger_structure = finger_data.get(
                    "finger_structure",
                    ""
                )

            palm_sheet.append([
                analysis.analysis_id,
                analysis.user_id,
                palm_shape,
                finger_structure,
                (
                    analysis.created_at.isoformat()
                    if analysis.created_at
                    else ""
                )
            ])

        # -----------------------------------------
        # TAROT REPORTS SHEET
        # -----------------------------------------

        tarot_sheet = workbook.create_sheet(
            "Tarot Reports"
        )

        tarot_sheet.append([
            "Reading ID",
            "User ID",
            "Spread",
            "Spread Type",
            "Card Count",
            "Date"
        ])

        for reading in tarot_readings:

            tarot_sheet.append([
                str(reading.reading_id),
                reading.user_id,
                reading.spread_name,
                reading.spread_type,
                reading.card_count,
                (
                    reading.created_at.isoformat()
                    if reading.created_at
                    else ""
                )
            ])

        # -----------------------------------------
        # PLATFORM ACTIVITY SHEET
        # -----------------------------------------

        activity_sheet = workbook.create_sheet(
            "Platform Activity"
        )

        activity_sheet.append([
            "Date",
            "Total Activity"
        ])

        activity_counter = Counter()

        for analysis in palm_analyses:

            if analysis.created_at:
                date = (
                    analysis.created_at
                    .date()
                    .isoformat()
                )

                activity_counter[date] += 1

        for reading in tarot_readings:

            if reading.created_at:
                date = (
                    reading.created_at
                    .date()
                    .isoformat()
                )

                activity_counter[date] += 1

        for date, count in sorted(
            activity_counter.items()
        ):

            activity_sheet.append([
                date,
                count
            ])

        # -----------------------------------------
        # SAVE FILE
        # -----------------------------------------

        reports_dir = Path("reports")
        reports_dir.mkdir(
            parents=True,
            exist_ok=True
        )

        output_path = (
            reports_dir /
            "system_reports.xlsx"
        )

        workbook.save(output_path)

        # -----------------------------------------
        # RETURN EXCEL FILE
        # -----------------------------------------

        return FileResponse(
            path=str(output_path),
            media_type=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
            filename="system_reports.xlsx"
        )

    except Exception as e:

        raise HTTPException(
            status_code=500,
            detail=(
                "Unable to export system reports: "
                f"{str(e)}"
            )
        )