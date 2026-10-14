import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_pptx(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    DARK_BG = RGBColor(15, 23, 42)      # Deep Slate #0f172a
    CARD_BG = RGBColor(30, 41, 59)      # Slate #1e293b
    ACCENT_GOLD = RGBColor(245, 158, 11) # Gold/Amber #f59e0b
    ACCENT_BLUE = RGBColor(14, 165, 233) # Sky Blue #0ea5e9
    ACCENT_GREEN = RGBColor(16, 185, 129)# Emerald #10b981
    TEXT_LIGHT = RGBColor(248, 250, 252) # White slate #f8fafc
    TEXT_MUTED = RGBColor(148, 163, 184) # Muted slate #94a3b8

    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, title_text, category_text):
        # Category tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf = cat_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category_text.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_GOLD

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_LIGHT

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, DARK_BG)

    # Title Card shape
    card1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.2), Inches(11.333), Inches(5.1))
    card1.fill.solid()
    card1.fill.fore_color.rgb = CARD_BG
    card1.line.color.rgb = ACCENT_BLUE
    card1.line.width = Pt(2)

    tf1 = card1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.6)
    tf1.margin_top = Inches(0.6)
    tf1.margin_right = Inches(0.6)

    p = tf1.paragraphs[0]
    p.text = "🔮 SOVEREIGN INTELLIGENCE PLATFORM"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GOLD
    p.space_after = Pt(14)

    p2 = tf1.add_paragraph()
    p2.text = "AI Based Palmistry & Tarot Intelligence Platform"
    p2.font.size = Pt(32)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_LIGHT
    p2.space_after = Pt(14)

    p3 = tf1.add_paragraph()
    p3.text = "Combining Computer Vision Palmistry, 78-Card Tarot Divination & Machine Learning Career Guidance"
    p3.font.size = Pt(16)
    p3.font.color.rgb = ACCENT_BLUE
    p3.space_after = Pt(28)

    p4 = tf1.add_paragraph()
    p4.text = "Executive Overview & 3 Core Operational Pillars Presentation"
    p4.font.size = Pt(14)
    p4.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 2: Executive Summary (3 Pillars)
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, DARK_BG)
    add_header(slide2, "Executive Summary: 3 Core Operational Pillars", "Platform Architecture")

    pillar_data = [
        ("🧠 Pillar 1: AI Interpretation Engine", "OPERATIONAL", ACCENT_BLUE, [
            "• Computer Vision Palm Line & Mount Extraction",
            "• Full 78-Card Tarot Deck (Celtic Cross & Spreads)",
            "• Natal Astrology & Multi-Modality Insight Scoring"
        ]),
        ("💡 Pillar 2: Recommendation System", "FUNCTIONAL", ACCENT_GOLD, [
            "• Personalized Daily Habit Rituals & Focus Sprints",
            "• Zodiac Crystal & Gemstone Remedies",
            "• 7-Chakra Energetic Alignment Protocols"
        ]),
        ("🧭 Pillar 3: Personalized Guidance", "COMPLETED", ACCENT_GREEN, [
            "• 🚀 ML 4-Phase Career Execution Roadmaps",
            "• 🧘 4 Somatic Stress-Relief Techniques",
            "• 🫂 1-on-1 Best-Friend Loneliness Chatbot"
        ])
    ]

    for i, (p_title, status, color, items) in enumerate(pillar_data):
        left = Inches(0.8 + i * 3.9)
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(3.6), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = color
        card.line.width = Pt(2)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.4)
        tf.margin_right = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = p_title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(10)

        p_stat = tf.add_paragraph()
        p_stat.text = f"STATUS: {status}"
        p_stat.font.size = Pt(11)
        p_stat.font.bold = True
        p_stat.font.color.rgb = color
        p_stat.space_after = Pt(16)

        for item in items:
            p_item = tf.add_paragraph()
            p_item.text = item
            p_item.font.size = Pt(12)
            p_item.font.color.rgb = TEXT_MUTED
            p_item.space_after = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 3: Pillar 1 - AI Interpretation Engine
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, DARK_BG)
    add_header(slide3, "Pillar 1: AI Interpretation Engine Details", "Core Intelligence 1")

    p1_boxes = [
        ("✋ Computer Vision Palmistry", [
            "• Detects 5 Major Lines: Life, Head, Heart, Fate & Sun",
            "• Maps 7 Planetary Mounts (Jupiter, Sun, Moon, Venus...)",
            "• Morphological edge gradients & 4 hand elements",
            "• Calculates line vigor confidence gauge (up to 98%)"
        ]),
        ("🎴 78-Card Tarot Sanctuary", [
            "• Full 22 Major Arcana + 56 Minor Arcana deck",
            "• 3D card shuffling & flip animation mechanics",
            "• 6 Divination Spreads (Celtic Cross, Life Path, 3-Card)",
            "• Clear interpretations and practical guidance notes"
        ]),
        ("🌌 Multi-Modality Scoring", [
            "• Parses Natal Zodiac charts & elemental balances",
            "• Synthesizes Palm + Tarot + Astrology data",
            "• Formula: 0.30(Palm) + 0.25(Tarot) + 0.20(Personality)...",
            "• API Endpoint: /api/readings/synthesize"
        ])
    ]

    for i, (b_title, b_items) in enumerate(p1_boxes):
        left = Inches(0.8 + i * 3.9)
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(3.6), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = ACCENT_BLUE
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.4)

        p = tf.paragraphs[0]
        p.text = b_title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE
        p.space_after = Pt(14)

        for item in b_items:
            pi = tf.add_paragraph()
            pi.text = item
            pi.font.size = Pt(12)
            pi.font.color.rgb = TEXT_MUTED
            pi.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 4: Pillar 2 - Recommendation System
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, DARK_BG)
    add_header(slide4, "Pillar 2: Personalized Recommendation System", "Core Intelligence 2")

    p2_boxes = [
        ("⚡ Daily Habit Rituals", [
            "• Morning Solar Focus Routine (Confidence booster)",
            "• Heart Line Open Dialogue Communication Practice",
            "• 25-minute Pomodoro Work/Study Sprints",
            "• Restorative Evening Offline Digital Detox"
        ]),
        ("💎 Crystal & Chakra Remedies", [
            "• Zodiac Crystal Matching (Amethyst, Rose Quartz, Citrine)",
            "• 7-Chakra energetic balancing exercises",
            "• Aura cleansing & daily positive affirmations",
            "• Element balancing (Fire, Water, Air, Earth)"
        ]),
        ("🎯 80/20 Vocation Guidance", [
            "• Personal growth friction auditing",
            "• Wealth discipline & financial focus habits",
            "• Leadership quotient boosting exercises",
            "• API Endpoint: /api/recommendations/daily"
        ])
    ]

    for i, (b_title, b_items) in enumerate(p2_boxes):
        left = Inches(0.8 + i * 3.9)
        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(3.6), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = ACCENT_GOLD
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.4)

        p = tf.paragraphs[0]
        p.text = b_title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = ACCENT_GOLD
        p.space_after = Pt(14)

        for item in b_items:
            pi = tf.add_paragraph()
            pi.text = item
            pi.font.size = Pt(12)
            pi.font.color.rgb = TEXT_MUTED
            pi.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 5: Pillar 3 - Career & Emotional Support
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, DARK_BG)
    add_header(slide5, "Pillar 3: Career Guidance & Emotional Support", "Core Intelligence 3")

    p3_grid = [
        ("🚀 ML Career Execution Engine", "Generates tailored 4-phase roadmaps (30-day, 90-day, 6-month, 1-year) for ANY career field (Developer, Marketer, Data Analyst, Doctor, Cyber Security)."),
        ("🧘 4 Somatic Stress-Relief Tools", "Science-backed somatic practices: Physiological Sigh (cortisol drop), 5-4-3-2-1 Grounding, Progressive Muscle Relaxation (PMR), and 4-7-8 Vagal Breathing."),
        ("🫂 Best-Friend Loneliness Chatbot", "Warm 1-on-1 companion protocol ('I'm right here with you, my friend! ❤️'). Provides comforting perspective shifts and 3 concrete steps to move forward."),
        ("💖 Uplifting Self-Care Rituals", "4 creative rituals: Micro-Pleasure Reset, Nature & Light Therapy, Emotional Release & Creative Journaling, and 3 Tiny Victories self-appreciation.")
    ]

    for i, (g_title, g_desc) in enumerate(p3_grid):
        row = i // 2
        col = i % 2
        left = Inches(0.8 + col * 5.9)
        top = Inches(1.8 + row * 2.5)

        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.6), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = ACCENT_GREEN
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = g_title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = ACCENT_GREEN
        p.space_after = Pt(8)

        pi = tf.add_paragraph()
        pi.text = g_desc
        pi.font.size = Pt(12)
        pi.font.color.rgb = TEXT_LIGHT

    # -------------------------------------------------------------
    # SLIDE 6: 4 User Roles & Dashboards
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, DARK_BG)
    add_header(slide6, "User Roles & RBAC Dashboard Architecture", "Platform Authorization")

    roles = [
        ("👤 Seeker (User)", "Aria Solstice", [
            "• Personal reading logs",
            "• Active spiritual goals",
            "• Daily guidance streak counter"
        ]),
        ("🎴 Tarot Reader", "Madame Seraphina", [
            "• Client consultation queue",
            "• Draw frequency heatmaps",
            "• Custom reader notes"
        ]),
        ("🧘 Spiritual Consultant", "Zachary Vance", [
            "• Seeker cohort monitoring",
            "• 30-day life trend projections",
            "• Growth velocity tracking"
        ]),
        ("⚙️ Administrator", "Grand Sovereign Admin", [
            "• Platform health monitoring",
            "• User RBAC permissions",
            "• Audit logs & security"
        ])
    ]

    for i, (r_role, r_persona, r_items) in enumerate(roles):
        left = Inches(0.8 + i * 2.9)
        card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(2.7), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = ACCENT_GOLD
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = r_role
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(4)

        pp = tf.add_paragraph()
        pp.text = f"Persona: {r_persona}"
        pp.font.size = Pt(10)
        pp.font.color.rgb = ACCENT_GOLD
        pp.space_after = Pt(12)

        for item in r_items:
            pi = tf.add_paragraph()
            pi.text = item
            pi.font.size = Pt(11)
            pi.font.color.rgb = TEXT_MUTED
            pi.space_after = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 7: Technology Stack
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, DARK_BG)
    add_header(slide7, "Technology Stack & Production Deployment", "System Architecture")

    tech_boxes = [
        ("⚙️ Backend & Machine Learning", [
            "• FastAPI Python 3.12 Framework",
            "• Uvicorn High-Performance Async Server",
            "• Scikit-Learn ML Models (RandomForest, KNN)",
            "• Google Gemini 2.5 Flash API SDK"
        ]),
        ("🎨 Frontend UI & 3D Visuals", [
            "• Tailwind CSS Glassmorphic Interface",
            "• Three.js 3D Particle Canvas Animation",
            "• Chart.js Radar & Trend Charts",
            "• Interactive 3D Tarot Card Flip Arcade"
        ]),
        ("🔒 Security & Export System", [
            "• JWT Token Authentication & Login Gate",
            "• Role-Based Access Control (RBAC)",
            "• ReportLab PDF Report Generation",
            "• openpyxl Excel Workbook Exporter"
        ])
    ]

    for i, (t_title, t_items) in enumerate(tech_boxes):
        left = Inches(0.8 + i * 3.9)
        card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(3.6), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = ACCENT_BLUE
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.4)

        p = tf.paragraphs[0]
        p.text = t_title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE
        p.space_after = Pt(14)

        for item in t_items:
            pi = tf.add_paragraph()
            pi.text = item
            pi.font.size = Pt(12)
            pi.font.color.rgb = TEXT_MUTED
            pi.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 8: Summary & Live Demo Links
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, DARK_BG)
    add_header(slide8, "Project Completion & Live Demo Access", "Production Release")

    card8 = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.8), Inches(11.333), Inches(4.8))
    card8.fill.solid()
    card8.fill.fore_color.rgb = CARD_BG
    card8.line.color.rgb = ACCENT_GREEN
    card8.line.width = Pt(2)

    tf8 = card8.text_frame
    tf8.word_wrap = True
    tf8.margin_left = Inches(0.5)
    tf8.margin_top = Inches(0.4)

    p = tf8.paragraphs[0]
    p.text = "🎉 All Project Requirements Successfully Fulfilled!"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(16)

    bullets = [
        "✔ AI Interpretation Engine: Synthesizing CV Palmistry, 78 Tarot Cards & Natal Astrology.",
        "✔ Recommendation System: Personalized daily habit rituals, crystal remedies & chakra balancing.",
        "✔ Personalized Guidance: 4-phase ML career roadmaps, 4 somatic stress tools & 1-on-1 best-friend chat.",
        "✔ Platform Security: Role-Based Access Control (RBAC) across 4 seed account personas.",
        "🌐 Live Web Application: http://127.0.0.1:8000",
        "🌐 Interactive Slide Deck: http://127.0.0.1:8000/static/presentation_deck.html",
        "🌐 OpenAPI Documentation: http://127.0.0.1:8000/docs"
    ]

    for b in bullets:
        pb = tf8.add_paragraph()
        pb.text = b
        pb.font.size = Pt(14)
        pb.font.color.rgb = TEXT_LIGHT
        pb.space_after = Pt(8)

    prs.save(output_path)
    try:
        print(f"[+] PPTX successfully generated at: {output_path}")
    except Exception:
        pass

if __name__ == "__main__":
    out_file = sys.argv[1] if len(sys.argv) > 1 else "Palmistry_Tarot_Platform_Presentation.pptx"
    build_pptx(out_file)
